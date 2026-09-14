"""Excel and PDF export helpers with the shared gold theme."""
import io

GOLD = "FFD54F"
INK = "1F2937"
BAND = "FFF8E1"


def excel_response(title, headers, rows, filename, totals_row=None):
    """Build a styled single-sheet workbook and return a StreamingResponse."""
    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True, color=INK)
        cell.fill = PatternFill("solid", fgColor=GOLD)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        ws.append(list(row))
    if totals_row:
        ws.append([])
        ws.append(list(totals_row))
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=BAND)
    for idx, header in enumerate(headers, start=1):
        widest = max([len(str(header))] +
                     [len(str(r[idx - 1])) for r in rows if idx <= len(r)] or [0])
        ws.column_dimensions[get_column_letter(idx)].width = min(40, max(11, widest + 3))
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"})


def pdf_response(title, headers, rows, filename, totals_row=None, landscape_mode=True):
    """Build a styled PDF table and return a StreamingResponse."""
    from fastapi.responses import StreamingResponse
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4) if landscape_mode else A4,
                            leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm)
    styles = getSampleStyleSheet()
    data = [list(headers)] + [list(r) for r in rows]
    if totals_row:
        data.append(list(totals_row))
    table = Table(data, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + GOLD)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#" + INK)),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E5E7EB")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#" + BAND)]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    if totals_row:
        style.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
        style.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#FFECB3")))
    table.setStyle(TableStyle(style))
    doc.build([Paragraph(title, styles["Title"]), Spacer(1, 6), table])
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename={filename}"})
