"""Drivers/Employees module: CRUD, documents, salary adjustments, exports."""
import io

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse

from services import sheets_service as ss
from utils.permissions import require_admin, require_admin_or_viewer
from utils.templates import render

router = APIRouter(tags=["drivers"])

ADJUSTMENT_TYPES = ("Advance", "Meals", "Incentive", "Deduction")


@router.get("/drivers")
async def drivers_page(request: Request, user=Depends(require_admin_or_viewer)):
    return render(request, "drivers.html")


@router.get("/drivers/{driver_id}")
async def driver_details_page(request: Request, driver_id: str,
                              user=Depends(require_admin_or_viewer)):
    driver = ss.find_record_by_id("drivers", driver_id)
    if not driver:
        raise HTTPException(status_code=404, detail="Driver not found")
    return render(request, "driver_details.html", {"driver": driver})


@router.get("/salaries")
async def salaries_page(request: Request, user=Depends(require_admin_or_viewer)):
    return render(request, "salaries.html")


@router.get("/api/drivers")
async def list_drivers(user=Depends(require_admin_or_viewer)):
    return sorted(ss.get_all_records("drivers"), key=lambda d: d.get("driver_name") or "")


@router.get("/api/drivers/{driver_id}")
async def get_driver(driver_id: str, user=Depends(require_admin_or_viewer)):
    driver = ss.find_record_by_id("drivers", driver_id)
    if not driver:
        raise HTTPException(status_code=404, detail="Driver not found")
    return driver


@router.post("/api/drivers")
async def add_driver(request: Request, user=Depends(require_admin)):
    data = await request.json()
    if not (data.get("driver_name") or "").strip():
        raise HTTPException(status_code=400, detail="Name is required")
    data.setdefault("status", "Active")
    data.setdefault("employee_type", "Driver")
    driver_id = ss.add_record("drivers", data)
    _sync_employee_access(data)
    ss.add_audit_log("Create", "Drivers", driver_id,
                     f"Added {data.get('employee_type')} {data.get('driver_name')}", user["email"])
    return {"success": True, "id": driver_id}


@router.put("/api/drivers/{driver_id}")
async def update_driver(driver_id: str, request: Request, user=Depends(require_admin)):
    data = await request.json()
    if not ss.find_record_by_id("drivers", driver_id):
        raise HTTPException(status_code=404, detail="Driver not found")
    ss.update_record("drivers", driver_id, data)
    _sync_employee_access(data)
    ss.add_audit_log("Update", "Drivers", driver_id,
                     f"Updated {data.get('driver_name', driver_id)}", user["email"])
    return {"success": True}


@router.delete("/api/drivers/{driver_id}")
async def delete_driver(driver_id: str, user=Depends(require_admin)):
    driver = ss.find_record_by_id("drivers", driver_id)
    if not driver:
        raise HTTPException(status_code=404, detail="Driver not found")
    ss.delete_record("drivers", driver_id)
    ss.add_audit_log("Delete", "Drivers", driver_id,
                     f"Deleted {driver.get('driver_name')}", user["email"])
    return {"success": True}


@router.get("/api/drivers/{driver_id}/documents")
async def driver_documents(driver_id: str, user=Depends(require_admin_or_viewer)):
    return [d for d in ss.get_all_records("documents")
            if d.get("entity_type") == "driver" and d.get("entity_id") == driver_id]


@router.post("/api/drivers/{driver_id}/documents")
async def upload_driver_document(driver_id: str, document_type: str = Form(...),
                                 file: UploadFile = File(...), user=Depends(require_admin)):
    from routes.vehicles import upload_entity_document
    return await upload_entity_document("driver", driver_id, document_type, file, user)


# ------------------------------------------------------- salary adjustments

@router.get("/api/salaries")
async def salary_summary(month: str = "", user=Depends(require_admin_or_viewer)):
    """Per-employee salary summary for a month: salary + incentives - advance/meals/deductions."""
    month = month or ss.this_month()
    adjustments = [a for a in ss.get_all_records("incentives") if a.get("for_month") == month]
    summary = []
    for driver in ss.get_all_records("drivers"):
        if (driver.get("status") or "Active") != "Active":
            continue
        rows = [a for a in adjustments if a.get("driver_id") == driver.get("driver_id")]
        totals = {t: sum(ss.to_float(a.get("amount")) for a in rows if a.get("entry_type") == t)
                  for t in ADJUSTMENT_TYPES}
        salary = ss.to_float(driver.get("salary"))
        net = salary + totals["Incentive"] - totals["Advance"] - totals["Meals"] - totals["Deduction"]
        summary.append({
            "driver_id": driver.get("driver_id"),
            "driver_name": driver.get("driver_name"),
            "employee_type": driver.get("employee_type"),
            "salary": salary,
            "advance": totals["Advance"],
            "meals": totals["Meals"],
            "incentive": totals["Incentive"],
            "deduction": totals["Deduction"],
            "net_payable": net,
        })
    return {"month": month, "rows": summary,
            "total_net": sum(r["net_payable"] for r in summary)}


@router.get("/api/salaries/adjustments")
async def list_adjustments(driver_id: str = "", month: str = "",
                           user=Depends(require_admin_or_viewer)):
    rows = ss.get_all_records("incentives")
    if driver_id:
        rows = [r for r in rows if r.get("driver_id") == driver_id]
    if month:
        rows = [r for r in rows if r.get("for_month") == month]
    return sorted(rows, key=lambda r: r.get("created_date") or "", reverse=True)


@router.post("/api/salaries/adjustments")
async def add_adjustment(request: Request, user=Depends(require_admin)):
    data = await request.json()
    if data.get("entry_type") not in ADJUSTMENT_TYPES:
        raise HTTPException(status_code=400, detail="Invalid adjustment type")
    driver = ss.find_record_by_id("drivers", data.get("driver_id", ""))
    if not driver:
        raise HTTPException(status_code=404, detail="Driver not found")
    data["driver_name"] = driver.get("driver_name")
    data.setdefault("for_month", ss.this_month())
    entry_id = ss.add_record("incentives", data)
    ss.add_audit_log("Create", "Salaries", entry_id,
                     f"{data['entry_type']} ₹{data.get('amount')} for {driver.get('driver_name')}",
                     user["email"])
    return {"success": True, "id": entry_id}


@router.delete("/api/salaries/adjustments/{entry_id}")
async def delete_adjustment(entry_id: str, user=Depends(require_admin)):
    ss.delete_record("incentives", entry_id)
    ss.add_audit_log("Delete", "Salaries", entry_id, "Deleted salary adjustment", user["email"])
    return {"success": True}


@router.get("/api/salaries/export")
async def export_salaries(month: str = "", user=Depends(require_admin_or_viewer)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    data = await salary_summary(month=month, user=user)
    wb = Workbook()
    ws = wb.active
    ws.title = f"Salaries {data['month']}"
    headers = ["Employee", "Type", "Salary", "Advance", "Meals",
               "Incentive", "Deduction", "Net Payable"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="1F2937")
        cell.fill = PatternFill("solid", fgColor="FFD54F")
    for r in data["rows"]:
        ws.append([r["driver_name"], r["employee_type"], r["salary"], r["advance"],
                   r["meals"], r["incentive"], r["deduction"], r["net_payable"]])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=salaries_{data['month']}.xlsx"})


def _sync_employee_access(data):
    """Auto-create an app user entry for employees that have an email."""
    email = (data.get("email") or "").strip().lower()
    if not email:
        return
    from services.auth_service import get_user_by_email
    if not get_user_by_email(email):
        ss.add_record("users", {
            "email": email, "name": data.get("driver_name", ""),
            "role": "driver", "status": "Active",
        })
