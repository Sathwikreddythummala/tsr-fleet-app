"""Expenses module: CRUD, filters, duplicate check, Excel/PDF export, purse sync."""
import io

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse

from services import sheets_service as ss
from utils import duplicate_check
from utils.permissions import require_editor, require_view
from utils.templates import render

router = APIRouter(tags=["expenses"])

EXPENSE_TYPES = ["Vehicle Expense", "Driver Expense", "House Expense", "Office Expense", "Other"]
CATEGORIES = ["Fuel", "Driver Expense", "Tyres", "Maintenance", "EMI", "Fastag",
              "Insurance", "Permit", "Accident", "Penalty", "Helper Expense",
              "Office Expense", "House Expense", "Other"]
HOUSE_SUBCATEGORIES = ["Rent", "Electricity", "Groceries", "Repairs", "Water",
                       "Family Expense", "Medical", "School/College Fees",
                       "Function Expense", "Other"]

# per-category subcategories for a fleet/vehicle maintenance business
SUBCATEGORIES = {
    "Fuel": ["Diesel", "Petrol", "CNG", "AdBlue", "Other"],
    "Driver Expense": ["Daily Bata", "Food", "Accommodation", "Advance",
                       "Medical", "Mobile Recharge", "Other"],
    "Tyres": ["New Tyre", "Retreading", "Puncture Repair", "Wheel Alignment",
              "Wheel Balancing", "Tube/Flap", "Other"],
    "Maintenance": ["General Service", "Engine Oil Change", "Oil/Air Filter",
                    "Engine Repair", "Clutch Work", "Brake Work",
                    "Gearbox/Transmission", "Suspension/Spring Work",
                    "Electrical Work", "Battery", "Radiator/Cooling",
                    "AC Repair", "Body Work/Denting", "Painting",
                    "Windshield/Glass", "Greasing", "Welding Work",
                    "Spare Parts", "Other"],
    "EMI": ["Vehicle EMI", "Other EMI"],
    "Fastag": ["Recharge", "New Tag", "Other"],
    "Insurance": ["New Policy", "Renewal", "Claim Expense", "Other"],
    "Permit": ["National Permit", "State Permit", "Permit Renewal", "Road Tax",
               "Fitness Fee", "PUC", "RTO Charges", "Other"],
    "Accident": ["Accident Repair", "Towing", "Third Party Payment",
                 "Legal/FIR Expense", "Other"],
    "Penalty": ["Traffic Challan", "Overload Fine", "RTO Fine", "Other"],
    "Helper Expense": ["Salary", "Food", "Advance", "Other"],
    "Office Expense": ["Rent", "Electricity", "Internet/Phone", "Stationery",
                       "Software", "Staff Welfare", "Travel", "Other"],
    "House Expense": HOUSE_SUBCATEGORIES,
    "Other": ["Other"],
}

PURSE_HOLDERS = ("TSR", "TRR")


@router.get("/expenses")
async def expenses_page(request: Request, user=Depends(require_view)):
    return render(request, "expenses.html", {
        "expense_types": EXPENSE_TYPES, "categories": CATEGORIES,
        "subcategories": SUBCATEGORIES,
    })


@router.get("/api/expenses")
async def list_expenses(date_from: str = "", date_to: str = "", vehicle: str = "",
                        category: str = "", sub_category: str = "", search: str = "",
                        month: str = "", page: int = 1, page_size: int = 25,
                        user=Depends(require_view)):
    rows = _filtered(date_from, date_to, vehicle, category, sub_category, search, month)
    total_amount = sum(ss.to_float(e.get("amount")) for e in rows)
    start = max(0, (page - 1) * page_size)
    return {
        "rows": rows[start:start + page_size],
        "total": len(rows),
        "page": page,
        "page_size": page_size,
        "total_amount": total_amount,
    }


@router.post("/api/expenses")
async def add_expense(request: Request, user=Depends(require_editor)):
    data = await request.json()
    if ss.to_float(data.get("amount")) <= 0:
        raise HTTPException(status_code=400, detail="Amount must be greater than zero")
    if duplicate_check.check("expenses", data):
        raise HTTPException(status_code=400, detail="Duplicate expense entry detected")
    data.setdefault("expense_date", ss.today_str())
    data.setdefault("for_month", (data.get("expense_date") or "")[:7] or ss.this_month())
    expense_id = ss.add_record("expenses", data)
    _sync_purse(expense_id, data)
    ss.add_audit_log("Create", "Expenses", expense_id,
                     f"{data.get('category')} ₹{data.get('amount')}", user["email"])
    return {"success": True, "id": expense_id}


@router.put("/api/expenses/{expense_id}")
async def update_expense(expense_id: str, request: Request, user=Depends(require_editor)):
    data = await request.json()
    if not ss.find_record_by_id("expenses", expense_id):
        raise HTTPException(status_code=404, detail="Expense not found")
    if duplicate_check.check("expenses", data, exclude_id=expense_id):
        raise HTTPException(status_code=400, detail="Duplicate expense entry detected")
    ss.update_record("expenses", expense_id, data)
    _remove_purse_link(expense_id)
    _sync_purse(expense_id, data)
    ss.add_audit_log("Update", "Expenses", expense_id, "Updated expense", user["email"])
    return {"success": True}


@router.delete("/api/expenses/{expense_id}")
async def delete_expense(expense_id: str, user=Depends(require_editor)):
    if not ss.find_record_by_id("expenses", expense_id):
        raise HTTPException(status_code=404, detail="Expense not found")
    ss.delete_record("expenses", expense_id)
    _remove_purse_link(expense_id)
    ss.add_audit_log("Delete", "Expenses", expense_id, "Deleted expense", user["email"])
    return {"success": True}


@router.get("/api/expenses/export/excel")
async def export_excel(date_from: str = "", date_to: str = "", vehicle: str = "",
                       category: str = "", sub_category: str = "", search: str = "",
                       month: str = "", user=Depends(require_view)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    rows = _filtered(date_from, date_to, vehicle, category, sub_category, search, month)
    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"
    headers = ["Date", "Month", "For", "Vehicle", "Driver", "Category",
               "Sub Category", "Description", "Amount", "Payment Mode", "Paid By"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="1F2937")
        cell.fill = PatternFill("solid", fgColor="FFD54F")
    for e in rows:
        ws.append([e.get("expense_date"), e.get("for_month"), e.get("expense_for"),
                   e.get("vehicle_number"), e.get("driver_name"), e.get("category"),
                   e.get("sub_category"), e.get("description"),
                   ss.to_float(e.get("amount")), e.get("payment_mode"), e.get("paid_by")])
    ws.append([])
    ws.append(["", "", "", "", "", "", "", "TOTAL",
               sum(ss.to_float(e.get("amount")) for e in rows)])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=expenses.xlsx"})


@router.get("/api/expenses/export/pdf")
async def export_pdf(date_from: str = "", date_to: str = "", vehicle: str = "",
                     category: str = "", sub_category: str = "", search: str = "",
                     month: str = "", user=Depends(require_view)):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    rows = _filtered(date_from, date_to, vehicle, category, sub_category, search, month)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=10 * mm, rightMargin=10 * mm)
    styles = getSampleStyleSheet()
    data = [["Date", "Vehicle", "Driver", "Category", "Sub Category",
             "Description", "Amount", "Paid By"]]
    for e in rows:
        data.append([e.get("expense_date"), e.get("vehicle_number"), e.get("driver_name"),
                     e.get("category"), e.get("sub_category"),
                     (e.get("description") or "")[:40],
                     f"{ss.to_float(e.get('amount')):,.2f}", e.get("paid_by")])
    data.append(["", "", "", "", "", "TOTAL",
                 f"{sum(ss.to_float(e.get('amount')) for e in rows):,.2f}", ""])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#FFD54F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E5E7EB")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#FFF8E1")]),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
    ]))
    doc.build([Paragraph("TSR Enterprises - Expense Report", styles["Title"]),
               Spacer(1, 6), table])
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": "attachment; filename=expenses.pdf"})


def _split(value):
    """Parse a comma-separated multi-select filter value into a list."""
    return [v.strip() for v in (value or "").split(",") if v.strip()]


def _filtered(date_from, date_to, vehicle, category, sub_category, search, month=""):
    rows = ss.get_all_records("expenses")
    months = _split(month)
    if months:
        rows = [e for e in rows
                if any((e.get("expense_date") or "").startswith(m) for m in months)]
    if date_from:
        rows = [e for e in rows if (e.get("expense_date") or "") >= date_from]
    if date_to:
        rows = [e for e in rows if (e.get("expense_date") or "") <= date_to]
    vehicles = _split(vehicle)
    if vehicles:
        rows = [e for e in rows if e.get("vehicle_number") in vehicles]
    categories = _split(category)
    if categories:
        rows = [e for e in rows if e.get("category") in categories]
    sub_categories = _split(sub_category)
    if sub_categories:
        rows = [e for e in rows if e.get("sub_category") in sub_categories]
    if search:
        q = search.lower()
        rows = [e for e in rows if q in " ".join(
            str(e.get(f) or "") for f in
            ("description", "vehicle_number", "driver_name", "category",
             "sub_category", "paid_by")).lower()]
    return sorted(rows, key=lambda e: e.get("expense_date") or "", reverse=True)


def _sync_purse(expense_id, data):
    """Expenses paid by a purse holder appear as purse Expense entries."""
    if (data.get("paid_by") or "").upper() in PURSE_HOLDERS:
        ss.add_record("purse", {
            "date": data.get("expense_date") or ss.today_str(),
            "holder": data["paid_by"].upper(),
            "type": "Expense",
            "amount": data.get("amount"),
            "description": f"{data.get('category')} - {data.get('description') or ''}",
            "vehicle_number": data.get("vehicle_number", ""),
            "category": data.get("category", ""),
            "reference_id": expense_id,
        })


def _remove_purse_link(expense_id):
    for p in ss.get_all_records("purse"):
        if p.get("reference_id") == expense_id:
            ss.delete_record("purse", p["purse_id"])
